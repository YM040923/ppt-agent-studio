import json
import re
import tempfile
import unittest
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation

from app.core.renderer import TemplateRenderer
from app.core.slide_planner import SlidePlanner
from app.service import GenerationService


HANDMADE_LAM = Path.home() / "Downloads" / "肺淋巴管平滑肌瘤病病因，表现.pptx"
TEMPLATE_DIR = (
    Path.home()
    / "Documents"
    / "Codex"
    / "2026-05-30"
    / "import-anygen"
    / "work"
    / "ppt-agent-data"
    / "templates"
    / "3bd9d6d273b54c63a961231036ea5695"
)


def _clean(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def _pptx_text_metrics(path: Path) -> dict:
    prs = Presentation(str(path))
    values = []
    for slide in prs.slides:
        total = 0
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                total += len(_clean(shape.text))
        values.append(total)
    return {
        "slides": len(values),
        "total": sum(values),
        "avg": round(sum(values) / max(1, len(values)), 1),
        "sparse_lt45": sum(value < 45 for value in values),
        "dense_ge160": sum(value >= 160 for value in values),
        "dense_ge220": sum(value >= 220 for value in values),
        "values": values,
    }


def _visible_slide_parts(path: Path) -> list[str]:
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with zipfile.ZipFile(path) as zf:
        pres_root = ET.fromstring(zf.read("ppt/presentation.xml"))
        rel_root = ET.fromstring(zf.read("ppt/_rels/presentation.xml.rels"))
        rid_to_target = {
            rel.attrib.get("Id"): rel.attrib.get("Target")
            for rel in rel_root
            if str(rel.attrib.get("Type") or "").endswith("/slide")
        }
        visible = []
        for child in pres_root:
            if not child.tag.endswith("sldIdLst"):
                continue
            for node in child:
                target = rid_to_target.get(node.attrib.get(f"{rel_ns}id"))
                if target:
                    visible.append(f"ppt/{target}" if not target.startswith("ppt/") else target)
        return visible


def _slide_xml_parts(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as zf:
        return [
            name
            for name in zf.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        ]


def _plan_text_units(slide: dict) -> int:
    chunks = [
        str(slide.get("title") or ""),
        str(slide.get("subtitle") or ""),
        *[str(item or "") for item in slide.get("toc_items", []) if isinstance(slide.get("toc_items"), list)],
        *[str(item or "") for item in slide.get("bullets", []) if isinstance(slide.get("bullets"), list)],
        *[
            str(item or "")
            for item in slide.get("summary_items", [])
            if isinstance(slide.get("summary_items"), list)
        ],
    ]
    if isinstance(slide.get("points"), list):
        for point in slide["points"]:
            if isinstance(point, dict):
                chunks.append(str(point.get("label") or ""))
                chunks.append(str(point.get("body") or ""))
    return sum(len(_clean(chunk)) for chunk in chunks)


@unittest.skipUnless(HANDMADE_LAM.exists() and (TEMPLATE_DIR / "template.pptx").exists(), "local LAM benchmark assets missing")
class HandmadeLamBenchmarkTests(unittest.TestCase):
    def test_reference_plan_matches_handmade_structure_and_density(self) -> None:
        handmade = _pptx_text_metrics(HANDMADE_LAM)
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        slides = plan["slides"]
        values = [_plan_text_units(slide) for slide in slides]
        titles = [str(slide.get("title") or "") for slide in slides]

        self.assertEqual(len(slides), handmade["slides"])
        self.assertGreaterEqual(sum(values), handmade["total"])
        self.assertGreaterEqual(sum(value >= 160 for value in values), handmade["dense_ge160"])
        self.assertLessEqual(sum(value < 45 for value in values), handmade["sparse_lt45"])
        self.assertIn("总结：LAM", titles[-2])
        self.assertIn("参考文献", titles[-1])
        self.assertTrue(any("病例线索" in title for title in titles[:5]))
        self.assertTrue(any("小杨第一幕" in title for title in titles))
        self.assertTrue(any("小杨第二幕" in title for title in titles))

    def test_rendered_reference_plan_retains_handmade_level_text(self) -> None:
        handmade = _pptx_text_metrics(HANDMADE_LAM)
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            generated = _pptx_text_metrics(output)

        self.assertLessEqual(generated["slides"], handmade["slides"] + 1)
        self.assertGreaterEqual(generated["total"], int(handmade["total"] * 0.99))
        self.assertGreaterEqual(generated["dense_ge160"], handmade["dense_ge160"] - 1)
        self.assertLessEqual(generated["sparse_lt45"], handmade["sparse_lt45"] + 1)
        self.assertGreaterEqual(generated["values"][-1], 500)

    def test_rendered_reference_plan_has_no_orphan_template_slides(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)

            self.assertEqual(len(_visible_slide_parts(output)), len(routed["slides"]))
            self.assertEqual(set(_slide_xml_parts(output)), set(_visible_slide_parts(output)))

    def test_rendered_reference_plan_does_not_leave_truncated_sentence_fragments(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            visible_text = "\n".join(
                shape.text
                for slide in prs.slides
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text
            )

        bad_fragments = [
            "LAM几乎全部发生于",
            "复发率高，胸膜固定等预防措施应比普通气胸更",
            "气胸可反复发生，且复",
            "临床上不应把LAM简",
        ]
        lines = [line.strip() for line in visible_text.splitlines() if line.strip()]
        for fragment in bad_fragments:
            self.assertNotIn(fragment, lines)

    def test_rendered_reference_plan_does_not_overpack_content_pages(self) -> None:
        handmade = _pptx_text_metrics(HANDMADE_LAM)
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            generated = _pptx_text_metrics(output)

        content_values = [
            value
            for slide, value in zip(routed["slides"], generated["values"])
            if slide.get("prototype_hint") == "content" and "参考" not in str(slide.get("title") or "")
        ]
        handmade_non_reference = handmade["values"][:-1]

        self.assertLessEqual(max(content_values), max(handmade_non_reference) + 65)
        self.assertLessEqual(sum(value > 280 for value in content_values), 2)
        self.assertLessEqual(
            sum(content_values) / len(content_values),
            (sum(value for value in handmade_non_reference if value >= 120) / len([value for value in handmade_non_reference if value >= 120])) + 55,
        )

    def test_rendered_reference_plan_has_no_underfilled_content_pages(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            generated = _pptx_text_metrics(output)

        content_values = [
            value
            for slide, value in zip(routed["slides"], generated["values"])
            if slide.get("prototype_hint") == "content" and "参考" not in str(slide.get("title") or "")
        ]
        self.assertGreaterEqual(min(content_values), 120)
        self.assertLessEqual(sum(value < 150 for value in content_values), 1)

    def test_rendered_multi_point_content_pages_are_not_single_text_blobs(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        blob_titles = []
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            for slide, planned in zip(prs.slides, routed["slides"]):
                if planned.get("prototype_hint") != "content":
                    continue
                title = str(planned.get("title") or "")
                if "参考" in title:
                    continue
                points = planned.get("points") if isinstance(planned.get("points"), list) else []
                bullets = planned.get("bullets") if isinstance(planned.get("bullets"), list) else []
                if len(points) < 3 and len(bullets) < 3:
                    continue
                visible_shapes = [
                    _clean(shape.text)
                    for shape in slide.shapes
                    if (
                        hasattr(shape, "text")
                        and shape.text
                        and shape.text.strip()
                        and shape.width > 0
                        and shape.height > 0
                        and shape.left < prs.slide_width
                        and shape.top < prs.slide_height
                    )
                ]
                substantial_shapes = [text for text in visible_shapes if len(text) >= 18]
                if len(substantial_shapes) <= 1 and sum(len(text) for text in visible_shapes) >= 160:
                    blob_titles.append(title)

        self.assertEqual(blob_titles, [])

    def test_rendered_reference_plan_preserves_medical_numeric_evidence(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            text = "\n".join(
                shape.text
                for slide in prs.slides
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text
            )

        self.assertIn("20岁女性", text)
        self.assertIn("80%", text)

    def test_reference_plan_does_not_overuse_one_content_template(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        content_routes = [
            int(slide.get("prototype_index") or 0)
            for slide in routed["slides"]
            if slide.get("prototype_hint") == "content" and "参考" not in str(slide.get("title") or "")
        ]
        route_counts = Counter(content_routes)

        self.assertGreaterEqual(len(route_counts), 3)
        self.assertLessEqual(max(route_counts.values()), 8)

    def test_reference_plan_uses_richer_content_layout_mix(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        content_routes = [
            int(slide.get("prototype_index") or 0)
            for slide in routed["slides"]
            if slide.get("prototype_hint") == "content" and "参考" not in str(slide.get("title") or "")
        ]
        route_counts = Counter(content_routes)

        self.assertGreaterEqual(len(route_counts), 5)
        self.assertTrue({4, 5, 6, 7}.intersection(route_counts))
        self.assertLessEqual(max(route_counts.values()), 6)

    def test_reference_plan_uses_dense_visual_routes_for_mechanism_or_treatment(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        spec["ai_standardization"] = GenerationService._local_template_standardization(spec)
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        target_routes = [
            int(slide.get("prototype_index") or 0)
            for slide in routed["slides"]
            if any(token in str(slide.get("title") or "") for token in ("mTOR", "治疗", "随访"))
        ]

        self.assertTrue(set(target_routes).intersection({8, 11, 12, 14}))

    def test_reference_plan_with_refreshed_local_contract_does_not_overuse_general_page(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        spec["ai_standardization"] = GenerationService._local_template_standardization(spec)
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        content_routes = [
            int(slide.get("prototype_index") or 0)
            for slide in routed["slides"]
            if slide.get("prototype_hint") == "content" and "参考" not in str(slide.get("title") or "")
        ]
        route_counts = Counter(content_routes)

        self.assertGreaterEqual(len(route_counts), 5)
        self.assertLessEqual(route_counts[14], 7)

    def test_rendered_dense_visual_mechanism_slide_keeps_body_detail(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        spec["ai_standardization"] = GenerationService._local_template_standardization(spec)
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            text = "\n".join(
                shape.text
                for slide, planned in zip(prs.slides, routed["slides"])
                if "mTOR通路异常" in str(planned.get("title") or "")
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text
            )

        self.assertIn("TSC1", text)
        self.assertIn("西罗莫司", text)
        self.assertGreaterEqual(len(_clean(text)), 160)

    def test_rendered_content_pages_keep_visible_page_titles(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        spec["ai_standardization"] = GenerationService._local_template_standardization(spec)
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        missing_titles = []
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            for slide, planned in zip(prs.slides, routed["slides"]):
                if planned.get("prototype_hint") != "content":
                    continue
                title = _clean(str(planned.get("title") or ""))
                if not title or "参考" in title:
                    continue
                visible_text = "".join(
                    _clean(shape.text)
                    for shape in slide.shapes
                    if (
                        hasattr(shape, "text")
                        and shape.text
                        and shape.text.strip()
                        and shape.width > 0
                        and shape.height > 0
                        and shape.left < prs.slide_width
                        and shape.top < prs.slide_height
                    )
                )
                if title[:8] not in visible_text:
                    missing_titles.append(str(planned.get("title") or ""))

        self.assertEqual(missing_titles, [])

    def test_rendered_content_pages_do_not_duplicate_title_into_body_when_title_is_visible(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        spec["ai_standardization"] = GenerationService._local_template_standardization(spec)
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        duplicated_titles = []
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            for slide, planned in zip(prs.slides, routed["slides"]):
                if planned.get("prototype_hint") != "content":
                    continue
                title = _clean(str(planned.get("title") or ""))
                if not title or "参考" in title:
                    continue
                visible_shapes = [
                    _clean(shape.text)
                    for shape in slide.shapes
                    if (
                        hasattr(shape, "text")
                        and shape.text
                        and shape.text.strip()
                        and shape.width > 0
                        and shape.height > 0
                        and shape.left < prs.slide_width
                        and shape.top < prs.slide_height
                    )
                ]
                if sum(1 for text in visible_shapes if title[:8] in text) > 1:
                    duplicated_titles.append(str(planned.get("title") or ""))

        self.assertEqual(duplicated_titles, [])

    def test_two_arrow_lam_type_slide_renders_both_sides(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            text = "\n".join(
                shape.text
                for slide, planned in zip(prs.slides, routed["slides"])
                if "散发性与TSC" in str(planned.get("title") or "")
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text
            )

        self.assertIn("散发性LAM", text)
        self.assertIn("TSC-LAM", text)
        self.assertGreaterEqual(len(_clean(text)), 140)

    def test_reference_plan_uses_visual_transition_template_for_sections(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        section_routes = [
            int(slide.get("prototype_index") or 0)
            for slide in routed["slides"]
            if slide.get("prototype_hint") == "section"
        ]

        self.assertEqual(len(section_routes), 4)
        self.assertTrue(section_routes)
        self.assertEqual(set(section_routes), {3})

    def test_rendered_reference_plan_section_pages_show_section_numbers(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            section_texts = []
            for slide, planned in zip(prs.slides, routed["slides"]):
                if planned.get("prototype_hint") != "section":
                    continue
                section_texts.append(
                    "\n".join(
                        shape.text
                        for shape in slide.shapes
                        if hasattr(shape, "text") and shape.text and shape.text.strip()
                    )
                )

        self.assertEqual(len(section_texts), 4)
        for expected, text in zip(("1", "2", "3", "4"), section_texts):
            self.assertRegex(text, rf"(^|\D){expected}($|\D)")

    def test_rendered_reference_plan_section_pages_keep_transition_scale(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "generated.pptx"
            renderer.render(template_path=template_path, plan=routed, output_path=output, template_spec=spec)
            prs = Presentation(str(output))
            section_sizes = []
            for slide, planned in zip(prs.slides, routed["slides"]):
                if planned.get("prototype_hint") != "section":
                    continue
                slide_sizes = []
                for shape in slide.shapes:
                    if not hasattr(shape, "text") or not shape.text.strip():
                        continue
                    sizes = [
                        run.font.size.pt
                        for paragraph in shape.text_frame.paragraphs
                        for run in paragraph.runs
                        if run.font.size
                    ]
                    if sizes:
                        slide_sizes.append((shape.text.strip(), max(sizes)))
                section_sizes.append(slide_sizes)

        self.assertEqual(len(section_sizes), 4)
        for slide_sizes in section_sizes:
            number_size = max(size for text, size in slide_sizes if re.fullmatch(r"\d{1,2}", text))
            title_size = max(size for text, size in slide_sizes if not re.fullmatch(r"\d{1,2}", text))
            self.assertGreaterEqual(number_size, 48)
            self.assertGreaterEqual(title_size, 30)

    def test_point_heavy_reference_slides_do_not_use_two_paragraph_template(self) -> None:
        template_path = TEMPLATE_DIR / "template.pptx"
        spec = json.loads((TEMPLATE_DIR / "spec.json").read_text(encoding="utf-8"))
        renderer = TemplateRenderer()
        plan = SlidePlanner().build_reference_quality_plan(
            "肺淋巴管平滑肌瘤病（LAM）的病因、临床表现与治疗",
            audience="PBL 病例汇报",
            tone="临床教学",
        )
        template = {
            "spec": spec,
            "pptx_path": str(template_path),
        }
        copy_patterns = renderer.copy_pattern_summary(template_path=template_path, slide_count=len(plan["slides"]))
        routed = GenerationService._apply_template_routes(plan, template, copy_patterns, force=True)

        bad_titles = [
            str(slide.get("title") or "")
            for slide in routed["slides"]
            if (
                slide.get("prototype_hint") == "content"
                and int(slide.get("prototype_index") or 0) == 5
                and len(slide.get("points") if isinstance(slide.get("points"), list) else []) >= 3
            )
        ]

        self.assertEqual(bad_titles, [])
