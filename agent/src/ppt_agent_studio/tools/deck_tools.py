from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from ppt_agent_studio.deck.spec import DeckSpec, SlideSpec, deck_from_outline, unique_slide_id
from ppt_agent_studio.preview.html_renderer import render_preview_document
from ppt_agent_studio.tools.base import ToolRegistry, ToolResult
from ppt_agent_studio.tools.catalog import core_tool_definitions


_HEX_COLOR = re.compile(r"^#(?:[0-9a-fA-F]{3}|[0-9a-fA-F]{6}|[0-9a-fA-F]{8})$")


def build_default_registry() -> ToolRegistry:
    registry = ToolRegistry()
    definitions = {definition.name: definition for definition in core_tool_definitions()}
    registry.register(definitions["deck.create_from_outline"], create_deck_from_outline)
    registry.register(definitions["deck.update_deck"], update_deck)
    registry.register(definitions["deck.update_slide"], update_slide)
    registry.register(definitions["deck.add_slide"], add_slide)
    registry.register(definitions["deck.move_slide"], move_slide)
    registry.register(definitions["deck.remove_slide"], remove_slide)
    registry.register(definitions["preview.render_html"], render_preview_html)
    registry.register(definitions["pptx.export"], export_pptx)
    registry.register(definitions["research.collect_brief"], collect_research_brief)
    registry.register(definitions["design.apply_theme"], apply_theme)
    return registry


def create_deck_from_outline(arguments: dict[str, Any]) -> ToolResult:
    outline = arguments.get("outline")
    if not isinstance(outline, dict):
        raise ValueError("outline must be an object")
    deck_id = str(arguments.get("deck_id") or "").strip()
    if not deck_id:
        raise ValueError("deck_id is required")
    revision = int(arguments.get("revision") or 0)
    deck = deck_from_outline(outline, deck_id=deck_id, revision=revision)
    return ToolResult(payload={"deck": deck.to_dict()})


def update_deck(arguments: dict[str, Any]) -> ToolResult:
    deck = _deck_from_arguments(arguments)
    patch = arguments.get("patch")
    if not isinstance(patch, dict):
        raise ValueError("patch must be an object")
    title = str(patch.get("title") or deck.title)
    metadata = dict(deck.metadata)
    patch_metadata = patch.get("metadata")
    if isinstance(patch_metadata, dict):
        for key, value in patch_metadata.items():
            text = _metadata_patch_text(value)
            if text:
                metadata[str(key)] = text
    updated = DeckSpec(
        deck_id=deck.deck_id,
        title=title,
        revision=deck.revision + 1,
        slides=deck.slides,
        theme=deck.theme,
        metadata=metadata,
    )
    return ToolResult(payload={"deck": updated.to_dict()})


def update_slide(arguments: dict[str, Any]) -> ToolResult:
    deck = _deck_from_arguments(arguments)
    slide_id = str(arguments.get("slide_id") or "").strip()
    patch = arguments.get("patch")
    if not slide_id:
        raise ValueError("slide_id is required")
    if not isinstance(patch, dict):
        raise ValueError("patch must be an object")

    slides: list[SlideSpec] = []
    found = False
    for slide in deck.slides:
        if slide.slide_id != slide_id:
            slides.append(slide)
            continue
        found = True
        slides.append(
            SlideSpec(
                slide_id=slide.slide_id,
                title=str(patch.get("title") or slide.title),
                layout=str(patch.get("layout") or slide.layout),
                blocks=_slide_patch_blocks(patch, slide),
                speaker_notes=_slide_patch_speaker_notes(patch, slide),
            )
        )
    if not found:
        raise ValueError(f"slide not found: {slide_id}")
    return ToolResult(payload={"deck": _deck_with_slides(deck, slides).to_dict()})


def add_slide(arguments: dict[str, Any]) -> ToolResult:
    deck = _deck_from_arguments(arguments)
    raw_slide = arguments.get("slide")
    if not isinstance(raw_slide, dict):
        raise ValueError("slide must be an object")
    after_slide_id = str(arguments.get("after_slide_id") or "").strip()
    before_slide_id = str(arguments.get("before_slide_id") or "").strip()
    if before_slide_id and after_slide_id:
        raise ValueError("before_slide_id and after_slide_id cannot both be provided")

    slides = list(deck.slides)
    used_slide_ids = {slide.slide_id for slide in slides}
    new_slide = _slide_from_dict(raw_slide, len(slides) + 1, used_slide_ids)
    if before_slide_id:
        for index, slide in enumerate(slides):
            if slide.slide_id == before_slide_id:
                slides.insert(index, new_slide)
                break
        else:
            raise ValueError(f"slide not found: {before_slide_id}")
    elif after_slide_id:
        for index, slide in enumerate(slides):
            if slide.slide_id == after_slide_id:
                slides.insert(index + 1, new_slide)
                break
        else:
            raise ValueError(f"slide not found: {after_slide_id}")
    else:
        slides.append(new_slide)
    return ToolResult(payload={"deck": _deck_with_slides(deck, slides).to_dict()})


def move_slide(arguments: dict[str, Any]) -> ToolResult:
    deck = _deck_from_arguments(arguments)
    slide_id = str(arguments.get("slide_id") or "").strip()
    after_slide_id = str(arguments.get("after_slide_id") or "").strip()
    before_slide_id = str(arguments.get("before_slide_id") or "").strip()
    if not slide_id:
        raise ValueError("slide_id is required")
    if before_slide_id and after_slide_id:
        raise ValueError("before_slide_id and after_slide_id cannot both be provided")
    if not before_slide_id and not after_slide_id:
        raise ValueError("before_slide_id or after_slide_id is required")
    if before_slide_id == slide_id or after_slide_id == slide_id:
        raise ValueError("cannot move slide relative to itself")

    moving_slide = None
    remaining: list[SlideSpec] = []
    for slide in deck.slides:
        if slide.slide_id == slide_id:
            moving_slide = slide
        else:
            remaining.append(slide)
    if moving_slide is None:
        raise ValueError(f"slide not found: {slide_id}")

    target_id = before_slide_id or after_slide_id
    for index, slide in enumerate(remaining):
        if slide.slide_id != target_id:
            continue
        insert_index = index if before_slide_id else index + 1
        remaining.insert(insert_index, moving_slide)
        break
    else:
        raise ValueError(f"slide not found: {target_id}")
    return ToolResult(payload={"deck": _deck_with_slides(deck, remaining).to_dict()})


def remove_slide(arguments: dict[str, Any]) -> ToolResult:
    deck = _deck_from_arguments(arguments)
    slide_id = str(arguments.get("slide_id") or "").strip()
    if not slide_id:
        raise ValueError("slide_id is required")
    slides = [slide for slide in deck.slides if slide.slide_id != slide_id]
    if len(slides) == len(deck.slides):
        raise ValueError(f"slide not found: {slide_id}")
    return ToolResult(payload={"deck": _deck_with_slides(deck, slides).to_dict()})


def render_preview_html(arguments: dict[str, Any]) -> ToolResult:
    raw_deck = arguments.get("deck")
    if not isinstance(raw_deck, dict):
        raise ValueError("deck must be an object")
    deck = _deck_from_dict(raw_deck)
    return ToolResult(payload={"html": render_preview_document(deck)})


def export_pptx(arguments: dict[str, Any]) -> ToolResult:
    raw_deck = arguments.get("deck")
    if not isinstance(raw_deck, dict):
        raise ValueError("deck must be an object")
    output_path_text = str(arguments.get("output_path") or "").strip()
    output_path = Path(output_path_text).expanduser()
    if not output_path.name:
        raise ValueError("output_path is required")
    if output_path.is_dir():
        raise ValueError("output_path must be a file path")
    if output_path.parent.exists() and not output_path.parent.is_dir():
        raise ValueError("output_path parent must be a directory")

    deck = _deck_from_dict(raw_deck)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = deck.title
    audience = _metadata_text(deck, "audience")
    style = _metadata_text(deck, "style")
    if audience:
        presentation.core_properties.subject = audience
    if style:
        presentation.core_properties.keywords = style
    blank_layout = presentation.slide_layouts[6]

    for index, slide in enumerate(deck.slides, start=1):
        ppt_slide = presentation.slides.add_slide(blank_layout)
        slide_background = _theme_rgb(deck, "slide_background", "#FFFFFF")
        text_color = _theme_rgb(deck, "text", "#111827")
        accent_color = _theme_rgb(deck, "accent", "#2563EB")
        ppt_slide.background.fill.solid()
        ppt_slide.background.fill.fore_color.rgb = slide_background
        _add_accent_bar(ppt_slide, presentation.slide_width, accent_color)
        _add_textbox(
            ppt_slide,
            left=0.65,
            top=6.95,
            width=0.55,
            height=0.2,
            text=f"{index:02d}",
            font_size=10,
            bold=True,
            font_color=accent_color,
        )
        if slide.speaker_notes:
            ppt_slide.notes_slide.notes_text_frame.text = slide.speaker_notes
        _add_textbox(
            ppt_slide,
            left=0.65,
            top=0.45,
            width=11.8,
            height=0.75,
            text=slide.title,
            font_size=30,
            bold=True,
            font_color=text_color,
        )
        y = 1.45
        for block in slide.blocks:
            block_text = _block_text(block)
            if not block_text:
                continue
            if block.get("type") == "point":
                _add_point_textbox(
                    ppt_slide,
                    left=0.85,
                    top=y,
                    width=11.1,
                    height=0.5,
                    label=str(block.get("label") or "").strip(),
                    body=str(block.get("body") or "").strip(),
                    font_color=text_color,
                )
            else:
                _add_textbox(
                    ppt_slide,
                    left=0.85,
                    top=y,
                    width=11.1,
                    height=0.5,
                    text=block_text,
                    font_size=18,
                    bold=block.get("type") == "subtitle",
                    font_color=text_color,
                )
            y += 0.58

    presentation.save(output_path)
    return ToolResult(payload={"path": str(output_path), "slide_count": len(deck.slides)})


def collect_research_brief(arguments: dict[str, Any]) -> ToolResult:
    topic = str(arguments.get("topic") or "").strip()
    audience = str(arguments.get("audience") or "").strip()
    raw_constraints = arguments.get("constraints")
    constraints = [str(item) for item in raw_constraints] if isinstance(raw_constraints, list) else []
    if not topic:
        raise ValueError("topic is required")
    if not audience:
        raise ValueError("audience is required")
    return ToolResult(
        payload={
            "brief": {
                "topic": topic,
                "audience": audience,
                "constraints": constraints,
                "questions": [
                    "Clarify the decision the deck must support.",
                    "Identify the audience's current belief and desired shift.",
                    "List the proof points needed for executive confidence.",
                ],
                "source_suggestions": [
                    f"Recent reports, filings, and investor materials related to {topic}.",
                    f"Industry benchmarks and analyst research tailored to {audience}.",
                    "Internal performance, customer, financial, and operating metrics.",
                ],
                "evidence_needs": _research_evidence_needs(constraints),
            }
        }
    )


def apply_theme(arguments: dict[str, Any]) -> ToolResult:
    raw_deck = arguments.get("deck")
    theme = arguments.get("theme")
    if not isinstance(raw_deck, dict):
        raise ValueError("deck must be an object")
    if not isinstance(theme, dict):
        raise ValueError("theme must be an object")
    deck = _deck_from_dict(raw_deck)
    themed = DeckSpec(
        deck_id=deck.deck_id,
        title=deck.title,
        revision=deck.revision + 1,
        slides=deck.slides,
        theme=dict(theme),
        metadata=deck.metadata,
    )
    return ToolResult(payload={"deck": themed.to_dict()})


def _deck_from_dict(raw_deck: dict[str, Any]) -> DeckSpec:
    raw_slides = raw_deck.get("slides") if isinstance(raw_deck.get("slides"), list) else []
    slides: list[SlideSpec] = []
    used_slide_ids: set[str] = set()
    for index, raw_slide in enumerate(raw_slides, start=1):
        if not isinstance(raw_slide, dict):
            continue
        slide_id = unique_slide_id(raw_slide.get("slide_id"), index, used_slide_ids)
        used_slide_ids.add(slide_id)
        slides.append(
            SlideSpec(
                slide_id=slide_id,
                title=str(raw_slide.get("title") or f"Slide {index}"),
                layout=str(raw_slide.get("layout") or "content"),
                blocks=_slide_blocks_from_dict(raw_slide),
                speaker_notes=_slide_speaker_notes(raw_slide),
            )
        )
    return DeckSpec(
        deck_id=str(raw_deck.get("deck_id") or "deck"),
        title=str(raw_deck.get("title") or "Untitled Deck"),
        revision=int(raw_deck.get("revision") or 0),
        slides=slides,
        theme=raw_deck.get("theme") if isinstance(raw_deck.get("theme"), dict) else {},
        metadata=raw_deck.get("metadata") if isinstance(raw_deck.get("metadata"), dict) else {},
    )


def _deck_from_arguments(arguments: dict[str, Any]) -> DeckSpec:
    raw_deck = arguments.get("deck")
    if not isinstance(raw_deck, dict):
        raise ValueError("deck must be an object")
    return _deck_from_dict(raw_deck)


def _slide_from_dict(raw_slide: dict[str, Any], index: int, used_slide_ids: set[str] | None = None) -> SlideSpec:
    slide_id = unique_slide_id(raw_slide.get("slide_id"), index, used_slide_ids or set())
    return SlideSpec(
        slide_id=slide_id,
        title=str(raw_slide.get("title") or f"Slide {index}"),
        layout=str(raw_slide.get("layout") or "content"),
        blocks=_slide_blocks_from_dict(raw_slide),
        speaker_notes=_slide_speaker_notes(raw_slide),
    )


def _slide_blocks_from_dict(raw_slide: dict[str, Any]) -> list[dict[str, Any]]:
    raw_blocks = raw_slide.get("blocks")
    if isinstance(raw_blocks, list):
        return _normalized_blocks(raw_blocks)
    normalized = deck_from_outline({"slides": [raw_slide]}, deck_id="_slide")
    return normalized.slides[0].blocks if normalized.slides else []


def _slide_patch_blocks(patch: dict[str, Any], slide: SlideSpec) -> list[dict[str, Any]]:
    raw_blocks = patch.get("blocks")
    if isinstance(raw_blocks, list):
        return _normalized_blocks(raw_blocks)
    if _has_outline_content_fields(patch):
        return _slide_blocks_from_dict(patch)
    return slide.blocks


def _normalized_blocks(raw_blocks: list[Any]) -> list[dict[str, Any]]:
    return [_normalized_block(block) for block in raw_blocks if isinstance(block, dict)]


def _normalized_block(block: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(block)
    block_type = str(normalized.get("type") or "text")
    text = str(normalized.get("text") or "").strip()
    if block_type != "point" and not text:
        alias = _block_text_alias(normalized)
        if alias:
            normalized["text"] = alias
    return normalized


def _block_text_alias(block: dict[str, Any]) -> str:
    for key in ("content", "value", "body", "label", "title"):
        value = block.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _has_outline_content_fields(source: dict[str, Any]) -> bool:
    return any(
        key in source
        for key in ("subtitle", "content", "body", "key_message", "toc_items", "points", "bullets", "summary_items")
    )


def _slide_speaker_notes(raw_slide: dict[str, Any]) -> str:
    for key in ("speaker_notes", "notes", "talk_track"):
        value = raw_slide.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _slide_patch_speaker_notes(patch: dict[str, Any], slide: SlideSpec) -> str:
    if any(key in patch for key in ("speaker_notes", "notes", "talk_track")):
        return _slide_speaker_notes(patch)
    return slide.speaker_notes


def _metadata_patch_text(value: Any) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, dict):
        for key in ("text", "value", "content", "label", "title"):
            text = value.get(key)
            if isinstance(text, str) and text.strip():
                return text.strip()
    return ""


def _deck_with_slides(deck: DeckSpec, slides: list[SlideSpec]) -> DeckSpec:
    return DeckSpec(
        deck_id=deck.deck_id,
        title=deck.title,
        revision=deck.revision + 1,
        slides=slides,
        theme=deck.theme,
        metadata=deck.metadata,
    )


def _add_accent_bar(slide: Any, slide_width: Any, accent_color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_width, Inches(0.08))
    shape.fill.solid()
    shape.fill.fore_color.rgb = accent_color
    shape.line.fill.background()


def _add_textbox(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    font_size: int,
    bold: bool,
    font_color: RGBColor | None = None,
) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if font_color is not None:
        run.font.color.rgb = font_color


def _add_point_textbox(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    label: str,
    body: str,
    font_color: RGBColor | None = None,
) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    paragraph = shape.text_frame.paragraphs[0]
    if label:
        label_run = paragraph.add_run()
        label_run.text = label
        label_run.font.size = Pt(18)
        label_run.font.bold = True
        if font_color is not None:
            label_run.font.color.rgb = font_color
    if body:
        body_run = paragraph.add_run()
        body_run.text = f": {body}" if label else body
        body_run.font.size = Pt(18)
        body_run.font.bold = False
        if font_color is not None:
            body_run.font.color.rgb = font_color


def _block_text(block: dict[str, Any]) -> str:
    block_type = str(block.get("type") or "text")
    if block_type == "point":
        label = str(block.get("label") or "").strip()
        body = str(block.get("body") or "").strip()
        if label and body:
            return f"{label}: {body}"
        return label or body
    text = str(block.get("text") or "").strip()
    if text and block_type in {"bullet", "summary_item", "toc_item"}:
        return f"\u2022 {text}"
    return text


def _theme_rgb(deck: DeckSpec, key: str, fallback: str) -> RGBColor:
    raw_value = deck.theme.get(key)
    value = raw_value.strip() if isinstance(raw_value, str) else fallback
    if not _HEX_COLOR.fullmatch(value):
        value = fallback
    value = value.lstrip("#")
    if len(value) == 3:
        value = "".join(character * 2 for character in value)
    if len(value) == 8:
        value = value[:6]
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _metadata_text(deck: DeckSpec, key: str) -> str:
    value = deck.metadata.get(key)
    return value.strip() if isinstance(value, str) else ""


def _research_evidence_needs(constraints: list[str]) -> list[str]:
    needs = [
        "Market context and size of the opportunity.",
        "Current-state baseline, pain points, and root causes.",
        "Decision options with value, risk, timing, and ownership implications.",
    ]
    if constraints:
        needs.append(f"Evidence that satisfies constraints: {', '.join(constraints)}.")
    return needs
