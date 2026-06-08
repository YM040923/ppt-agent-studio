import asyncio

from pptx import Presentation

from ppt_agent_studio.runtime.session import AgentSession
from ppt_agent_studio.tools.base import ToolDefinition, ToolRegistry, ToolResult


class FakeOutlinePlanner:
    def __init__(self):
        self.prompts = []

    async def create_outline(self, prompt):
        self.prompts.append(prompt)
        return {
            "deck_title": "Planner Deck",
            "slides": [
                {
                    "title": "Planner Deck",
                    "subtitle": "Generated through an injected planner",
                    "prototype_hint": "cover",
                }
            ],
        }


class ThemedOutlinePlanner:
    async def create_outline(self, prompt):
        return {
            "deck_title": "Themed Deck",
            "theme": {
                "name": "executive-consulting",
                "background": "#EEF2F7",
                "slide_background": "#FFFFFF",
                "text": "#111827",
                "accent": "#2563EB",
            },
            "slides": [
                {
                    "title": "Themed Deck",
                    "subtitle": "Generated with a planner theme",
                    "prototype_hint": "cover",
                }
            ],
        }


def test_agent_session_turn_emits_ordered_preview_and_pptx_events(tmp_path):
    session = AgentSession(session_id="session_001", deck_id="deck_001", artifact_dir=tmp_path)

    async def collect():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    events = asyncio.run(collect())

    assert [event.type for event in events] == [
        "user.message",
        "plan.updated",
        "tool.completed",
        "plan.updated",
        "tool.completed",
        "deck.updated",
        "tool.completed",
        "preview.ready",
        "tool.completed",
        "pptx.ready",
        "plan.updated",
    ]
    assert [event.seq for event in events] == list(range(1, 12))
    assert events[1].payload["plan"]["plan_id"] == "deck_001-r1-plan"
    assert events[1].payload["plan"]["status"] == "running"
    assert events[1].payload["plan"]["slide_count"] == 5
    assert events[1].payload["plan"]["steps"][0]["status"] == "running"
    assert events[1].payload["plan"]["steps"][1]["status"] == "completed"
    assert events[1].payload["plan"]["steps"][2]["step_id"] == "draft_slides"
    assert events[1].payload["plan"]["steps"][2]["title"] == "Draft 5 slides"
    assert events[1].payload["plan"]["steps"][2]["status"] == "pending"
    assert events[2].payload == {
        "tool_name": "research.collect_brief",
        "status": "completed",
        "summary": "Collected research brief.",
    }
    assert events[3].payload["research_brief"]["audience"] == "general business audience"
    assert events[3].payload["plan"]["steps"][0]["status"] == "completed"
    assert events[3].payload["plan"]["steps"][2]["status"] == "running"
    assert events[4].payload == {
        "tool_name": "deck.create_from_outline",
        "status": "completed",
        "summary": "Created deck from outline.",
    }
    assert events[4].deck_revision == 1
    assert events[5].deck_revision == 1
    assert events[6].payload == {
        "tool_name": "preview.render_html",
        "status": "completed",
        "summary": "Rendered live preview HTML.",
    }
    assert events[7].payload["deck_id"] == "deck_001"
    assert events[7].payload["revision"] == 1
    assert events[7].payload["deck_title"] == "Board AI Strategy"
    assert events[7].payload["slide_count"] == 5
    assert events[7].payload["html"].startswith("<!doctype html>")
    assert "Board AI Strategy" in events[7].payload["html"]
    assert events[8].payload == {
        "tool_name": "pptx.export",
        "status": "completed",
        "summary": "Exported editable PPTX artifact.",
    }
    assert events[9].payload["path"].endswith("deck_001-r1.pptx")
    assert events[9].payload["slide_count"] == 5
    assert len(Presentation(events[9].payload["path"]).slides) == 5
    assert events[10].payload["plan"]["status"] == "completed"
    assert all(step["status"] == "completed" for step in events[10].payload["plan"]["steps"])


def test_agent_session_ignores_empty_user_message():
    session = AgentSession(session_id="session_002", deck_id="deck_002")

    async def collect():
        return [event async for event in session.submit_user_message("   ")]

    events = asyncio.run(collect())

    assert len(events) == 1
    assert events[0].type == "error"
    assert events[0].payload["message"] == "User message is required"


def test_agent_session_uses_injected_outline_planner(tmp_path):
    planner = FakeOutlinePlanner()
    session = AgentSession(
        session_id="session_003",
        deck_id="deck_003",
        outline_planner=planner,
        artifact_dir=tmp_path,
    )

    async def collect():
        return [event async for event in session.submit_user_message("Make a CFO margin story")]

    events = asyncio.run(collect())

    assert planner.prompts == ["Make a CFO margin story"]
    assert events[1].payload["outline"]["deck_title"] == "Planner Deck"
    assert events[5].payload["deck"]["title"] == "Planner Deck"
    assert "Planner Deck" in events[7].payload["html"]
    assert events[9].payload["slide_count"] == 1


def test_agent_session_preserves_planner_theme_for_preview_and_export(tmp_path):
    session = AgentSession(
        session_id="session_themed",
        deck_id="deck_themed",
        outline_planner=ThemedOutlinePlanner(),
        artifact_dir=tmp_path,
    )

    async def collect():
        return [event async for event in session.submit_user_message("Make an executive deck")]

    events = asyncio.run(collect())
    presentation = Presentation(events[9].payload["path"])

    assert events[5].payload["deck"]["theme"]["name"] == "executive-consulting"
    assert "--slide-accent: #2563EB;" in events[7].payload["html"]
    assert events[9].payload["deck_title"] == "Themed Deck"
    assert events[9].payload["theme_name"] == "executive-consulting"
    assert str(presentation.slides[0].shapes[0].fill.fore_color.rgb) == "2563EB"


def test_agent_session_preserves_prompt_metadata_for_preview_and_export(tmp_path):
    session = AgentSession(session_id="session_metadata", deck_id="deck_metadata", artifact_dir=tmp_path)

    async def collect():
        return [
            event
            async for event in session.submit_user_message(
                "Make a 3 page AI strategy deck for the executive committee, style McKinsey"
            )
        ]

    events = asyncio.run(collect())
    presentation = Presentation(events[9].payload["path"])

    assert events[3].payload["research_brief"]["audience"] == "executive committee"
    assert events[5].payload["deck"]["metadata"] == {
        "audience": "executive committee",
        "style": "McKinsey",
    }
    assert '<meta name="ppt-agent-audience" content="executive committee" />' in events[7].payload["html"]
    assert presentation.core_properties.subject == "executive committee"
    assert presentation.core_properties.keywords == "McKinsey"


def test_agent_session_creates_slide_on_follow_up_request(tmp_path):
    session = AgentSession(session_id="session_followup", deck_id="deck_followup", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [event async for event in session.submit_user_message("Create a risk mitigation slide")]

    asyncio.run(first_turn())
    events = asyncio.run(second_turn())

    assert [event.type for event in events] == [
        "user.message",
        "plan.updated",
        "tool.completed",
        "deck.updated",
        "tool.completed",
        "preview.ready",
        "tool.completed",
        "pptx.ready",
        "plan.updated",
    ]
    presentation = Presentation(events[7].payload["path"])
    assert [event.seq for event in events] == list(range(12, 21))
    assert events[2].payload == {
        "tool_name": "deck.add_slide",
        "status": "completed",
        "summary": "Added slide: Risk Mitigation.",
    }
    assert events[3].deck_revision == 2
    assert events[3].payload["deck"]["revision"] == 2
    assert events[3].payload["deck"]["slides"][-1]["title"] == "Risk Mitigation"
    assert events[5].payload["slide_count"] == 6
    assert events[7].payload["slide_count"] == 6
    assert len(presentation.slides) == 6
    assert events[8].payload["plan"]["status"] == "completed"


def test_agent_session_does_not_treat_address_as_add_slide_request(tmp_path):
    session = AgentSession(session_id="session_followup_words", deck_id="deck_followup_words", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [event async for event in session.submit_user_message("Address market sizing gaps")]

    asyncio.run(first_turn())
    events = asyncio.run(second_turn())

    assert events[2].payload["tool_name"] == "research.collect_brief"


def test_agent_session_removes_numbered_slide_on_follow_up_request(tmp_path):
    session = AgentSession(session_id="session_remove", deck_id="deck_remove", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [event async for event in session.submit_user_message("Remove slide 2")]

    first_events = asyncio.run(first_turn())
    original_second_title = first_events[5].payload["deck"]["slides"][1]["title"]
    events = asyncio.run(second_turn())

    assert [event.type for event in events] == [
        "user.message",
        "plan.updated",
        "tool.completed",
        "deck.updated",
        "tool.completed",
        "preview.ready",
        "tool.completed",
        "pptx.ready",
        "plan.updated",
    ]
    presentation = Presentation(events[7].payload["path"])
    assert events[2].payload == {
        "tool_name": "deck.remove_slide",
        "status": "completed",
        "summary": "Removed slide 2.",
    }
    assert events[3].payload["deck"]["revision"] == 2
    assert len(events[3].payload["deck"]["slides"]) == 4
    assert all(slide["title"] != original_second_title for slide in events[3].payload["deck"]["slides"])
    assert events[5].payload["slide_count"] == 4
    assert events[7].payload["slide_count"] == 4
    assert len(presentation.slides) == 4
    assert events[8].payload["plan"]["status"] == "completed"


def test_agent_session_removes_ordinal_slide_on_follow_up_request(tmp_path):
    session = AgentSession(session_id="session_remove_ordinal", deck_id="deck_remove_ordinal", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [event async for event in session.submit_user_message("Remove the second slide")]

    first_events = asyncio.run(first_turn())
    original_second_title = first_events[5].payload["deck"]["slides"][1]["title"]
    events = asyncio.run(second_turn())

    assert events[2].payload == {
        "tool_name": "deck.remove_slide",
        "status": "completed",
        "summary": "Removed slide 2.",
    }
    assert len(events[3].payload["deck"]["slides"]) == 4
    assert all(slide["title"] != original_second_title for slide in events[3].payload["deck"]["slides"])


def test_agent_session_rejects_out_of_range_slide_remove_follow_up(tmp_path):
    session = AgentSession(session_id="session_remove_range", deck_id="deck_remove_range", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [event async for event in session.submit_user_message("Remove slide 99")]

    first_events = asyncio.run(first_turn())
    events = asyncio.run(second_turn())

    assert [event.type for event in events] == ["user.message", "error"]
    assert events[1].payload == {"message": "Slide 99 is not available. Deck has 5 slides."}
    assert session.deck is not None
    assert session.deck.revision == first_events[5].payload["deck"]["revision"]
    assert len(session.deck.slides) == 5


def test_agent_session_renames_numbered_slide_on_follow_up_request(tmp_path):
    session = AgentSession(session_id="session_update", deck_id="deck_update", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [
            event
            async for event in session.submit_user_message(
                "Rename slide 2 to Executive AI Roadmap"
            )
        ]

    asyncio.run(first_turn())
    events = asyncio.run(second_turn())

    assert [event.type for event in events] == [
        "user.message",
        "plan.updated",
        "tool.completed",
        "deck.updated",
        "tool.completed",
        "preview.ready",
        "tool.completed",
        "pptx.ready",
        "plan.updated",
    ]
    presentation = Presentation(events[7].payload["path"])
    assert events[2].payload == {
        "tool_name": "deck.update_slide",
        "status": "completed",
        "summary": "Renamed slide 2: Executive AI Roadmap.",
    }
    assert events[3].payload["deck"]["revision"] == 2
    assert events[3].payload["deck"]["slides"][1]["title"] == "Executive AI Roadmap"
    assert "Executive AI Roadmap" in events[5].payload["html"]
    assert events[7].payload["slide_count"] == 5
    assert presentation.slides[1].shapes[2].text_frame.paragraphs[0].runs[0].text == "Executive AI Roadmap"
    assert events[8].payload["plan"]["status"] == "completed"


def test_agent_session_renames_ordinal_slide_on_follow_up_request(tmp_path):
    session = AgentSession(session_id="session_update_ordinal", deck_id="deck_update_ordinal", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [
            event
            async for event in session.submit_user_message(
                "Rename the second slide to Executive AI Roadmap"
            )
        ]

    asyncio.run(first_turn())
    events = asyncio.run(second_turn())

    assert events[2].payload == {
        "tool_name": "deck.update_slide",
        "status": "completed",
        "summary": "Renamed slide 2: Executive AI Roadmap.",
    }
    assert events[3].payload["deck"]["slides"][1]["title"] == "Executive AI Roadmap"


def test_agent_session_rejects_out_of_range_slide_rename_follow_up(tmp_path):
    session = AgentSession(session_id="session_update_range", deck_id="deck_update_range", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [event async for event in session.submit_user_message("Rename slide 99 to Appendix")]

    first_events = asyncio.run(first_turn())
    events = asyncio.run(second_turn())

    assert [event.type for event in events] == ["user.message", "error"]
    assert events[1].payload == {"message": "Slide 99 is not available. Deck has 5 slides."}
    assert session.deck is not None
    assert session.deck.revision == first_events[5].payload["deck"]["revision"]
    assert len(session.deck.slides) == 5


def test_agent_session_applies_dark_theme_on_plain_darker_follow_up_request(tmp_path):
    session = AgentSession(session_id="session_theme_followup", deck_id="deck_theme_followup", artifact_dir=tmp_path)

    async def first_turn():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def second_turn():
        return [event async for event in session.submit_user_message("Make it darker")]

    asyncio.run(first_turn())
    events = asyncio.run(second_turn())

    assert [event.type for event in events] == [
        "user.message",
        "plan.updated",
        "tool.completed",
        "deck.updated",
        "tool.completed",
        "preview.ready",
        "tool.completed",
        "pptx.ready",
        "plan.updated",
    ]
    presentation = Presentation(events[7].payload["path"])
    assert events[2].payload == {
        "tool_name": "design.apply_theme",
        "status": "completed",
        "summary": "Applied theme: executive-dark.",
    }
    assert events[3].payload["deck"]["revision"] == 2
    assert events[3].payload["deck"]["theme"] == {
        "name": "executive-dark",
        "background": "#0B1220",
        "slide_background": "#111827",
        "text": "#F9FAFB",
        "accent": "#38BDF8",
    }
    assert events[5].payload["theme_name"] == "executive-dark"
    assert "--slide-background: #111827;" in events[5].payload["html"]
    assert "--slide-text: #F9FAFB;" in events[5].payload["html"]
    assert events[7].payload["deck_title"] == "Board AI Strategy"
    assert events[7].payload["theme_name"] == "executive-dark"
    assert str(presentation.slides[0].shapes[0].fill.fore_color.rgb) == "38BDF8"
    assert events[8].payload["plan"]["status"] == "completed"


def test_agent_session_applies_light_theme_on_plain_brighter_follow_up_request(tmp_path):
    session = AgentSession(session_id="session_light_theme_followup", deck_id="deck_light_theme_followup", artifact_dir=tmp_path)

    async def create_deck():
        return [event async for event in session.submit_user_message("Make a 5 slide board AI strategy deck")]

    async def apply_dark_theme():
        return [event async for event in session.submit_user_message("Apply a dark executive theme")]

    async def apply_light_theme():
        return [event async for event in session.submit_user_message("Make it brighter")]

    asyncio.run(create_deck())
    asyncio.run(apply_dark_theme())
    events = asyncio.run(apply_light_theme())

    assert [event.type for event in events] == [
        "user.message",
        "plan.updated",
        "tool.completed",
        "deck.updated",
        "tool.completed",
        "preview.ready",
        "tool.completed",
        "pptx.ready",
        "plan.updated",
    ]
    presentation = Presentation(events[7].payload["path"])
    assert events[2].payload["tool_name"] == "design.apply_theme"
    assert events[3].payload["deck"]["revision"] == 3
    assert events[3].payload["deck"]["theme"] == {
        "name": "executive-consulting",
        "background": "#EEF2F7",
        "slide_background": "#FFFFFF",
        "text": "#111827",
        "accent": "#2563EB",
    }
    assert events[5].payload["theme_name"] == "executive-consulting"
    assert "--slide-background: #FFFFFF;" in events[5].payload["html"]
    assert str(presentation.slides[0].shapes[0].fill.fore_color.rgb) == "2563EB"


def test_agent_session_uses_tool_registry_for_deck_and_preview():
    calls = []
    registry = ToolRegistry()

    def create_deck(args):
        calls.append(("deck.create_from_outline", args["deck_id"]))
        return ToolResult(
            payload={
                "deck": {
                    "deck_id": args["deck_id"],
                    "title": "Custom Deck",
                    "revision": args["revision"],
                    "slides": [],
                }
            }
        )

    def render_preview(args):
        calls.append(("preview.render_html", args["deck"]["title"]))
        return ToolResult(payload={"html": "<!doctype html><title>Custom</title>"})

    registry.register(
        ToolDefinition(
            name="deck.create_from_outline",
            description="Create a deck.",
            input_schema={"type": "object"},
        ),
        create_deck,
    )
    registry.register(
        ToolDefinition(
            name="preview.render_html",
            description="Render preview HTML.",
            input_schema={"type": "object"},
        ),
        render_preview,
    )
    registry.register(
        ToolDefinition(
            name="research.collect_brief",
            description="Collect research brief.",
            input_schema={"type": "object"},
        ),
        lambda args: calls.append(("research.collect_brief", args["audience"])) or ToolResult(payload={"brief": args}),
    )
    registry.register(
        ToolDefinition(
            name="pptx.export",
            description="Export a deck.",
            input_schema={"type": "object"},
        ),
        lambda args: ToolResult(payload={"path": args["output_path"], "slide_count": 0}),
    )
    session = AgentSession(session_id="session_004", deck_id="deck_004", tool_registry=registry)

    async def collect():
        return [event async for event in session.submit_user_message("Make a sales deck")]

    events = asyncio.run(collect())

    assert calls == [
        ("research.collect_brief", "general business audience"),
        ("deck.create_from_outline", "deck_004"),
        ("preview.render_html", "Custom Deck"),
    ]
    assert events[5].payload["deck"]["title"] == "Custom Deck"
    assert events[7].payload["html"] == "<!doctype html><title>Custom</title>"
    assert events[9].payload["path"].endswith("deck_004-r1.pptx")
