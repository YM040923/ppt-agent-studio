import asyncio

from pptx import Presentation

from ppt_agent_studio.tools.deck_tools import build_default_registry


def test_pptx_export_tool_writes_editable_deck(tmp_path):
    registry = build_default_registry()
    deck = {
        "deck_id": "deck_001",
        "title": "AI Strategy",
        "revision": 1,
        "metadata": {
            "audience": "executive committee",
            "style": "McKinsey",
        },
        "slides": [
            {
                "slide_id": "s1",
                "title": "AI Strategy",
                "layout": "cover",
                "blocks": [{"type": "subtitle", "text": "Board briefing"}],
                "speaker_notes": "Open with the decision the board needs to make.",
            },
            {
                "slide_id": "s2",
                "title": "Priorities",
                "layout": "content",
                "blocks": [
                    {"type": "bullet", "text": "Focus the operating model."},
                    {"type": "point", "label": "Sequence", "body": "Ship in measurable waves."},
                ],
            },
        ],
    }
    output_path = tmp_path / "ai-strategy.pptx"

    async def run():
        return await registry.run("pptx.export", {"deck": deck, "output_path": str(output_path)})

    result = asyncio.run(run())
    presentation = Presentation(str(output_path))
    first_slide_texts = [shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text")]
    second_slide_texts = [shape.text for shape in presentation.slides[1].shapes if hasattr(shape, "text")]
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    assert result.payload["path"] == str(output_path)
    assert result.payload["slide_count"] == 2
    assert len(presentation.slides) == 2
    assert "01" in first_slide_texts
    assert "02" in second_slide_texts
    assert "AI Strategy" in text
    assert "Board briefing" in text
    assert "Focus the operating model." in text
    assert "Sequence" in text
    assert "Ship in measurable waves." in text
    point_shape = next(shape for shape in presentation.slides[1].shapes if hasattr(shape, "text") and "Sequence" in shape.text)
    point_runs = point_shape.text_frame.paragraphs[0].runs
    assert [run.text for run in point_runs] == ["Sequence", ": Ship in measurable waves."]
    assert point_runs[0].font.bold is True
    assert point_runs[1].font.bold is False
    assert "Open with the decision the board needs to make." in presentation.slides[0].notes_slide.notes_text_frame.text
    assert presentation.core_properties.subject == "executive committee"
    assert presentation.core_properties.keywords == "McKinsey"


def test_pptx_export_tool_applies_theme_to_editable_shapes(tmp_path):
    registry = build_default_registry()
    deck = {
        "deck_id": "deck_theme",
        "title": "Theme Test",
        "revision": 1,
        "theme": {
            "slide_background": "#111827",
            "text": "#F8FAFC",
            "accent": "#22C55E",
        },
        "slides": [
            {
                "slide_id": "s1",
                "title": "Theme Test",
                "layout": "cover",
                "blocks": [{"type": "subtitle", "text": "Executive readout"}],
            }
        ],
    }
    output_path = tmp_path / "theme-test.pptx"

    async def run():
        return await registry.run("pptx.export", {"deck": deck, "output_path": str(output_path)})

    asyncio.run(run())
    slide = Presentation(str(output_path)).slides[0]
    title_shape = next(shape for shape in slide.shapes if hasattr(shape, "text") and shape.text == "Theme Test")

    assert str(slide.background.fill.fore_color.rgb) == "111827"
    assert str(slide.shapes[0].fill.fore_color.rgb) == "22C55E"
    assert str(title_shape.text_frame.paragraphs[0].runs[0].font.color.rgb) == "F8FAFC"
